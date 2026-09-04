"""DOCX parser: Word paragraph styles map directly to block kinds. "Heading 1" starts a
new chapter; deeper heading levels stay inside the current chapter as heading blocks.
See PLAN.md 'ingest/base.py'.
"""
import re
from pathlib import Path

import docx as python_docx

from ..models import BlockKind
from .base import ParseError
from .document import DocBlock, DocChapter, Document

_HEADING_RE = re.compile(r"^heading\s+(\d+)$", re.IGNORECASE)
_QUOTE_STYLES = {"quote", "intense quote"}


def _classify(style_name: str) -> tuple[BlockKind, int | None]:
    """Returns (kind, heading_level). heading_level is None for non-headings."""
    m = _HEADING_RE.match(style_name or "")
    if m:
        return BlockKind.heading, int(m.group(1))
    if (style_name or "").strip().lower() in _QUOTE_STYLES:
        return BlockKind.quote, None
    return BlockKind.para, None


class DocxParser:
    extensions = ("docx",)

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower().lstrip(".") in self.extensions

    def parse(self, path: Path) -> Document:
        try:
            doc = python_docx.Document(str(path))
        except Exception as exc:
            raise ParseError(f"could not read DOCX {path}: {exc}") from exc

        chapters: list[DocChapter] = []
        current_blocks: list[DocBlock] = []
        current_title = "Chapter 1"
        chapter_index = 0
        seen_h1 = False

        def flush() -> None:
            nonlocal chapter_index, current_blocks, current_title
            if current_blocks:
                chapters.append(DocChapter(index=chapter_index, title=current_title, blocks=current_blocks))
                chapter_index += 1
            current_blocks = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            kind, level = _classify(para.style.name if para.style else "")
            if kind is BlockKind.heading and level == 1:
                flush()
                current_title = text if len(text) <= 80 else text[:77] + "..."
                seen_h1 = True
                continue
            current_blocks.append(DocBlock(kind=kind, text=text))
        flush()

        if not chapters:
            raise ParseError(f"{path} contains no text")
        if not seen_h1 and len(chapters) == 1:
            chapters[0].title = doc.core_properties.title or path.stem

        title = doc.core_properties.title or path.stem
        author = doc.core_properties.author or None
        language = (doc.core_properties.language or "und").split("-")[0].lower() or "und"

        return Document(title=title, author=author, language=language, chapters=chapters)
